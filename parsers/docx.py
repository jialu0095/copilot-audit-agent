from __future__ import annotations
from pathlib import Path
from docx import Document


async def extract_docx_text(file_path: str) -> str:
    p = Path(file_path)
    if not p.exists():
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    try:
        doc = Document(str(p))
        parts: list[str] = []

        for para in doc.paragraphs:
            t = (para.text or "").strip()
            if t:
                parts.append(t)

        # table content (control points, RACI, approval chains are often in tables)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    t = (cell.text or "").strip()
                    if t:
                        parts.append(t)

        return "\n".join(parts).strip()
    except Exception as e:
        raise RuntimeError(f"Failed to parse DOCX: {e}") from e